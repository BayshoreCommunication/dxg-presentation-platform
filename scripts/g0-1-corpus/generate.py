#!/usr/bin/env python3
"""Generate the G0-1 Room Agent test corpus (docs/PHASE0_GATE.md G0-1, docs/poc/G0-1_PLAN.md A5).

Writes every fixture into tests/fixtures/g0-1/ plus MANIFEST.md and manifest.json.

    scripts/g0-1-corpus/.venv/bin/python scripts/g0-1-corpus/generate.py            # committed corpus
    scripts/g0-1-corpus/.venv/bin/python scripts/g0-1-corpus/generate.py --large    # + large/ deck (~80 MB, gitignored)
    scripts/g0-1-corpus/.venv/bin/python scripts/g0-1-corpus/generate.py --keynote  # + .key via Keynote on macOS

Determinism: every .pptx/.pptm and ffmpeg output is byte-identical between runs (fixed timestamps,
fixed zip dates, bitexact encoders, seeded noise). Files produced by LibreOffice (.ppt, .pdf) and
Keynote (.key) embed their own timestamps and are content-identical only.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import posixpath
import random
import re
import shutil
import struct
import subprocess
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import PROG_ID
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
OUT = REPO / "tests" / "fixtures" / "g0-1"
LARGE = OUT / "large"
BUILD = HERE / ".build"  # intermediates (clips, posters, LibreOffice profile); gitignored

FFMPEG = shutil.which("ffmpeg") or "/opt/homebrew/bin/ffmpeg"
FFPROBE = shutil.which("ffprobe") or "/opt/homebrew/bin/ffprobe"
SOFFICE = shutil.which("soffice") or "/opt/homebrew/bin/soffice"

FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)
ZIP_TIME = (2026, 1, 1, 0, 0, 0)

W169, H169 = Inches(13.333), Inches(7.5)
W43, H43 = Inches(10), Inches(7.5)
W1610, H1610 = Inches(12), Inches(7.5)

P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
RT_VIDEO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video"
RT_MEDIA = "http://schemas.microsoft.com/office/2007/relationships/media"
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
RT_OLE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
RT_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
RT_HLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
RT_VBA = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
CT_PRES_MAIN = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
CT_PRES_MACRO = "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"
CT_VBA = "application/vnd.ms-office.vbaProject"

PALETTE = ["1F4E79", "2E7D32", "B71C1C", "6A1B9A", "E65100", "00695C", "37474F", "AD1457"]


# --------------------------------------------------------------------------------------------
# Manifest bookkeeping
# --------------------------------------------------------------------------------------------
@dataclass
class Entry:
    name: str
    category: str
    tests: str
    items: str
    expected: str
    how: str
    todo: str = ""
    corrupt: bool = False
    committed: bool = True


ENTRIES: list[Entry] = []


def register(**kw) -> None:
    ENTRIES.append(Entry(**kw))


def log(msg: str) -> None:
    print(msg, flush=True)


def run(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    res = subprocess.run(cmd, capture_output=True, text=True, **kw)
    if res.returncode != 0:
        raise RuntimeError(f"command failed ({res.returncode}): {' '.join(cmd)}\n{res.stderr[-2000:]}")
    return res


# --------------------------------------------------------------------------------------------
# Media clips (ffmpeg)
# --------------------------------------------------------------------------------------------
BITEXACT = ["-fflags", "+bitexact", "-flags:v", "+bitexact", "-flags:a", "+bitexact", "-map_metadata", "-1",
            "-threads", "1"]


def ff_clip(out: Path, *, size: str, seconds: int, vcodec: list[str], acodec: list[str], extra: list[str] = ()):
    if out.exists():
        return out
    cmd = [FFMPEG, "-y", "-hide_banner", "-loglevel", "error",
           "-f", "lavfi", "-i", f"testsrc2=size={size}:rate=25:duration={seconds}",
           "-f", "lavfi", "-i", f"sine=frequency=440:sample_rate=48000:duration={seconds}",
           *vcodec, *acodec, *BITEXACT, *extra, "-shortest", str(out)]
    run(cmd)
    return out


def ff_poster(clip: Path, out: Path) -> Path:
    if not out.exists():
        run([FFMPEG, "-y", "-hide_banner", "-loglevel", "error", "-ss", "1", "-i", str(clip),
             "-frames:v", "1", *BITEXACT, str(out)])
    return out


def build_media() -> dict[str, Path]:
    BUILD.mkdir(parents=True, exist_ok=True)
    h264 = ["-c:v", "libx264", "-preset", "medium", "-crf", "30", "-pix_fmt", "yuv420p", "-profile:v", "high"]
    aac = ["-c:a", "aac", "-b:a", "64k"]
    m = {
        "h264": ff_clip(BUILD / "clip-h264.mp4", size="480x270", seconds=4, vcodec=h264, acodec=aac,
                        extra=["-movflags", "+faststart"]),
        "hevc": ff_clip(BUILD / "clip-hevc.mp4", size="480x270", seconds=4,
                        vcodec=["-c:v", "libx265", "-preset", "medium", "-crf", "32", "-pix_fmt", "yuv420p",
                                "-tag:v", "hvc1", "-x265-params", "log-level=error:pools=none:frame-threads=1"],
                        acodec=aac, extra=["-movflags", "+faststart"]),
        "prores": ff_clip(BUILD / "clip-prores.mov", size="480x270", seconds=3,
                          vcodec=["-c:v", "prores_ks", "-profile:v", "0", "-pix_fmt", "yuv422p10le",
                                  "-vendor", "apl0"],
                          acodec=["-c:a", "pcm_s16le"]),
        "h264_1080": ff_clip(BUILD / "video-1080.mp4", size="1920x1080", seconds=10,
                             vcodec=["-c:v", "libx264", "-preset", "medium", "-crf", "32", "-pix_fmt", "yuv420p",
                                     "-profile:v", "high"], acodec=aac, extra=["-movflags", "+faststart"]),
        "hevc_1080": ff_clip(BUILD / "video-1080-hevc.mp4", size="1920x1080", seconds=10,
                             vcodec=["-c:v", "libx265", "-preset", "medium", "-crf", "34", "-pix_fmt", "yuv420p",
                                     "-tag:v", "hvc1",
                                     "-x265-params", "log-level=error:pools=none:frame-threads=1"],
                             acodec=aac, extra=["-movflags", "+faststart"]),
    }
    for k in ("h264", "hevc", "prores"):
        m[f"{k}_poster"] = ff_poster(m[k], BUILD / f"poster-{k}.png")
    m["still"] = BUILD / "still-1080.png"
    if not m["still"].exists():
        run([FFMPEG, "-y", "-hide_banner", "-loglevel", "error", "-f", "lavfi", "-i",
             "testsrc2=size=1920x1080:rate=1:duration=1", "-frames:v", "1", *BITEXACT, str(m["still"])])
    return m


# --------------------------------------------------------------------------------------------
# Slide building helpers
# --------------------------------------------------------------------------------------------
def new_prs(width=W169, height=H169, title="G0-1 corpus deck") -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = width, height
    cp = prs.core_properties
    cp.title = title
    cp.author = "DXG G0-1 corpus generator"
    cp.last_modified_by = "DXG G0-1 corpus generator"
    cp.created = cp.modified = cp.last_printed = FIXED_TIME
    cp.revision = 1
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, text, left, top, width, height, size=28, bold=False, font=None, color="FFFFFF",
            align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        if font:
            r.font.name = font
    return tb


def background(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hex_color)


def box(slide, text, left, top, width, height, color="FFC000", shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=24,
        font=None, text_color="000000"):
    shp = slide.shapes.add_shape(shape, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color)
    shp.line.fill.background()
    tf = shp.text_frame
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string(text_color)
            if font:
                r.font.name = font
    return shp


def content_slide(prs, deck: str, n: int, total: int, title: str, body: str = "", font=None):
    """A visually distinct slide: coloured background, big slide counter (for screenshot checks)."""
    s = blank(prs)
    W, H = prs.slide_width, prs.slide_height
    background(s, PALETTE[(n - 1) % len(PALETTE)])
    textbox(s, title, Inches(0.6), Inches(0.4), W - Inches(1.2), Inches(1.2), size=40, bold=True, font=font)
    if body:
        textbox(s, body, Inches(0.6), Inches(1.7), W - Inches(4.2), H - Inches(2.8), size=24, font=font)
    textbox(s, f"{n} / {total}", W - Inches(3.4), H - Inches(2.2), Inches(3.0), Inches(1.6), size=72, bold=True,
            font=font, align=PP_ALIGN.RIGHT)
    textbox(s, deck, Inches(0.6), H - Inches(0.8), W - Inches(4.2), Inches(0.5), size=14, color="DDDDDD")
    return s


_SLD_ORDER = ["cSld", "clrMapOvr", "transition", "timing", "extLst"]


def insert_sld_child(sld, el):
    """Insert el into p:sld respecting the schema order cSld, clrMapOvr, transition, timing, extLst."""
    local = etree.QName(el).localname
    for old in sld.findall(qn(f"p:{local}")):
        sld.remove(old)
    after = _SLD_ORDER[_SLD_ORDER.index(local) + 1:]
    for child in sld:
        if etree.QName(child).localname in after:
            child.addprevious(el)
            return el
    sld.append(el)
    return el


def set_transition(slide, inner_xml: str, spd="med", adv_tm: int | None = None, adv_click=True):
    attrs = f'spd="{spd}"'
    if not adv_click:
        attrs += ' advClick="0"'
    if adv_tm is not None:
        attrs += f' advTm="{adv_tm}"'
    el = parse_xml(f"<p:transition {nsdecls('p')} {attrs}>{inner_xml}</p:transition>")
    insert_sld_child(slide._element, el)


# --------------------------------------------------------------------------------------------
# Animation timing (p:timing) builder
# --------------------------------------------------------------------------------------------
EFFECT_DUR = {"appear": 1, "fade_in": 500, "fly_in": 500, "spin": 2000, "fade_out": 500, "path": 2000,
              "path_curve": 2000, "play_media": None}


class Ids:
    def __init__(self, start=1):
        self.n = start - 1

    def __call__(self):
        self.n += 1
        return self.n


def _tgt(spid):
    return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'


def _set_vis(ids, spid, val, delay=0):
    return (f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{delay}"/>'
            f'</p:stCondLst></p:cTn>{_tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
            f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="{val}"/></p:to></p:set>')


def _effect(ids, kind, spid, node_type, media_ms=0):
    """Return (<p:par> xml, duration_ms) for one effect."""
    ctn = ids()
    if kind == "appear":
        head = 'presetID="1" presetClass="entr" presetSubtype="0"'
        body = _set_vis(ids, spid, "visible")
    elif kind == "fade_in":
        head = 'presetID="10" presetClass="entr" presetSubtype="0"'
        body = (_set_vis(ids, spid, "visible") +
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{ids()}" dur="500"/>'
                f'{_tgt(spid)}</p:cBhvr></p:animEffect>')
    elif kind == "fly_in":  # Fly In, from bottom
        head = 'presetID="2" presetClass="entr" presetSubtype="4"'
        body = _set_vis(ids, spid, "visible")
        for attr, frm, to in (("ppt_x", "#ppt_x", "#ppt_x"), ("ppt_y", "1+#ppt_h/2", "#ppt_y")):
            body += (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base">'
                     f'<p:cTn id="{ids()}" dur="500" fill="hold"/>{_tgt(spid)}<p:attrNameLst>'
                     f'<p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                     f'<p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav>'
                     f'<p:tav tm="100000"><p:val><p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>')
    elif kind == "spin":
        head = 'presetID="8" presetClass="emph" presetSubtype="0"'
        body = (f'<p:animRot by="21600000"><p:cBhvr><p:cTn id="{ids()}" dur="2000" fill="hold"/>{_tgt(spid)}'
                f'<p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>')
    elif kind == "fade_out":
        head = 'presetID="10" presetClass="exit" presetSubtype="0"'
        body = (f'<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn id="{ids()}" dur="500"/>'
                f'{_tgt(spid)}</p:cBhvr></p:animEffect>' + _set_vis(ids, spid, "hidden", delay=499))
    elif kind in ("path", "path_curve"):
        head = 'presetID="0" presetClass="path" presetSubtype="0" accel="50000" decel="50000"'
        if kind == "path":  # straight line to the right, 40% of slide width
            path, pts, rctr = "M 0 0 L 0.4 0 E", "AA", '<p:rCtr x="20000" y="0"/>'
        else:  # S-curve down and right
            path = "M 0 0 C 0.15 -0.25 0.25 0.25 0.4 0.2 E"
            pts, rctr = "AA", '<p:rCtr x="20000" y="10000"/>'
        body = (f'<p:animMotion origin="layout" path="{path}" pathEditMode="relative" rAng="0" ptsTypes="{pts}">'
                f'<p:cBhvr><p:cTn id="{ids()}" dur="2000" fill="hold"/>{_tgt(spid)}<p:attrNameLst>'
                f'<p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
                f'{rctr}</p:animMotion>')
    elif kind == "play_media":
        head = 'presetID="1" presetClass="mediacall" presetSubtype="0"'
        body = (f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{ids()}" dur="{media_ms}" '
                f'fill="hold"/>{_tgt(spid)}</p:cBhvr></p:cmd>')
    else:
        raise ValueError(kind)
    grp = '' if kind == "play_media" else ' grpId="0"'
    xml = (f'<p:par><p:cTn id="{ctn}" {head} fill="hold"{grp} nodeType="{node_type}"><p:stCondLst>'
           f'<p:cond delay="0"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')
    dur = media_ms if kind == "play_media" else EFFECT_DUR[kind]
    return xml, dur


def set_timing(slide, groups, videos=(), media_ms=0):
    """Replace the slide's p:timing.

    groups: list of click groups. Each group is {"auto": bool, "steps": [[(kind, spid), ...], ...]}.
    Effects inside a step run together ("with previous"); each later step runs "after previous".
    The first effect of a group is a clickEffect (or afterEffect when auto=True, i.e. starts on slide load).
    videos: shape ids of movie shapes that need a p:video media node.
    """
    ids = Ids()
    root_id, seq_id = ids(), ids()
    anim_spids: list[int] = []
    groups_xml = ""
    for gi, g in enumerate(groups):
        gctn = ids()
        if g.get("auto"):
            assert gi == 0, "only the first group can auto-start"
            cond = f'<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="{seq_id}"/></p:cond>'
        else:
            cond = '<p:cond delay="indefinite"/>'
        steps_xml, t = "", 0
        for si, step in enumerate(g["steps"]):
            sctn = ids()
            effs, longest = "", 0
            for ei, (kind, spid) in enumerate(step):
                if ei > 0:
                    nt = "withEffect"
                elif si > 0 or g.get("auto"):
                    nt = "afterEffect"
                else:
                    nt = "clickEffect"
                x, d = _effect(ids, kind, spid, nt, media_ms)
                effs += x
                longest = max(longest, d)
                if kind != "play_media" and spid not in anim_spids:
                    anim_spids.append(spid)
            steps_xml += (f'<p:par><p:cTn id="{sctn}" fill="hold"><p:stCondLst><p:cond delay="{t}"/>'
                          f'</p:stCondLst><p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par>')
            t += longest
        groups_xml += (f'<p:par><p:cTn id="{gctn}" fill="hold"><p:stCondLst>{cond}</p:stCondLst>'
                       f'<p:childTnLst>{steps_xml}</p:childTnLst></p:cTn></p:par>')
    seq = ""
    if groups:
        seq = (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{seq_id}" dur="indefinite" nodeType="mainSeq">'
               f'<p:childTnLst>{groups_xml}</p:childTnLst></p:cTn>'
               '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
               '</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl>'
               '</p:cond></p:nextCondLst></p:seq>')
    vids = ""
    for spid in videos:
        vids += (f'<p:video><p:cMediaNode vol="80000"><p:cTn id="{ids()}" fill="hold" display="0"><p:stCondLst>'
                 f'<p:cond delay="indefinite"/></p:stCondLst></p:cTn>{_tgt(spid)}</p:cMediaNode></p:video>')
    bld = "".join(f'<p:bldP spid="{s}" grpId="0" animBg="1"/>' for s in anim_spids)
    bld = f"<p:bldLst>{bld}</p:bldLst>" if bld else ""
    xml = (f'<p:timing {nsdecls("p")}><p:tnLst><p:par><p:cTn id="{root_id}" dur="indefinite" restart="never" '
           f'nodeType="tmRoot"><p:childTnLst>{seq}{vids}</p:childTnLst></p:cTn></p:par></p:tnLst>{bld}</p:timing>')
    insert_sld_child(slide._element, parse_xml(xml))


# --------------------------------------------------------------------------------------------
# Package (zip) post-processing
# --------------------------------------------------------------------------------------------
def read_zip(path: Path) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(path) as z:
        return [(i.filename, z.read(i.filename)) for i in z.infolist()]


def write_zip(path: Path, entries: list[tuple[str, bytes]]) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, data)
    path.write_bytes(buf.getvalue())


def rels_path_for(part: str) -> str:
    d, f = posixpath.split(part)
    return posixpath.join(d, "_rels", f + ".rels")


def drop_unreferenced(entries):
    """Remove media/embedding parts no longer referenced by any internal relationship."""
    referenced = set()
    for name, data in entries:
        if not name.endswith(".rels"):
            continue
        src_dir = posixpath.dirname(posixpath.dirname(name))  # folder of the source part
        for rel in etree.fromstring(data):
            if rel.get("TargetMode") == "External":
                continue
            referenced.add(posixpath.normpath(posixpath.join(src_dir, rel.get("Target"))).lstrip("/"))
    dropped = {n for n, _ in entries if n.startswith(("ppt/media/", "ppt/embeddings/")) and n not in referenced}
    out = []
    for n, d in entries:
        if n in dropped:
            continue
        if n == "[Content_Types].xml":  # remove Overrides that point at dropped parts
            root = etree.fromstring(d)
            for ov in root.findall(f"{{{CT_NS}}}Override"):
                if ov.get("PartName").lstrip("/") in dropped:
                    root.remove(ov)
            d = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        out.append((n, d))
    return out


def save_pptx(prs, path: Path, mutate=None):
    """Save via python-pptx, apply optional package mutation, rewrite zip deterministically."""
    buf = io.BytesIO()
    prs.save(buf)
    tmp = BUILD / ("_tmp_" + path.name)
    tmp.write_bytes(buf.getvalue())
    entries = read_zip(tmp)
    tmp.unlink()
    if mutate:
        entries = mutate(entries)
    write_zip(path, entries)


def edit_entry(entries, name, fn):
    out, found = [], False
    for n, d in entries:
        if n == name:
            d, found = fn(d), True
        out.append((n, d))
    if not found:
        raise KeyError(name)
    return out


def externalise(entries, slide_part: str, rel_types: tuple[str, ...], target: str, new_type: str | None = None):
    """Point the slide's relationships of the given types at an external target."""
    def fix(data):
        root = etree.fromstring(data)
        hit = 0
        for rel in root:
            if rel.get("Type") in rel_types:
                rel.set("Target", target)
                rel.set("TargetMode", "External")
                if new_type:
                    rel.set("Type", new_type)
                hit += 1
        assert hit, f"no {rel_types} relationship in {slide_part}"
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return drop_unreferenced(edit_entry(entries, rels_path_for(slide_part), fix))


# --------------------------------------------------------------------------------------------
# Minimal OLE2 / Compound File Binary writer (for the structural vbaProject.bin)
# --------------------------------------------------------------------------------------------
def cfb_single_stream(stream_name: str, payload: bytes) -> bytes:
    """A valid CFB v3 file with one root storage and one regular (non-mini) stream."""
    SECT = 512
    FREE, END, FATSECT, NOSTREAM = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD, 0xFFFFFFFF
    if len(payload) < 4096:
        payload = payload + b" " * (4096 - len(payload))  # keep it out of the mini stream
    n_data = -(-len(payload) // SECT)
    assert 2 + n_data <= SECT // 4, "payload too large for a single FAT sector"
    header = bytearray(SECT)
    struct.pack_into("<8s16sHHHHH6sIIIIIIIII", header, 0, bytes.fromhex("D0CF11E0A1B11AE1"), b"\0" * 16,
                     0x003E, 0x0003, 0xFFFE, 9, 6, b"\0" * 6, 0, 1, 1, 0, 4096, END, 0, END, 0)
    struct.pack_into("<I", header, 76, 0)
    for i in range(1, 109):
        struct.pack_into("<I", header, 76 + 4 * i, FREE)
    fat = [FATSECT, END] + [3 + i for i in range(n_data - 1)] + [END]
    fat += [FREE] * (SECT // 4 - len(fat))
    fat_bytes = struct.pack(f"<{SECT // 4}I", *fat)

    def dirent(name, typ, child=NOSTREAM, start=END, size=0):
        e = bytearray(128)
        if name:
            enc = name.encode("utf-16-le") + b"\0\0"
            e[0:len(enc)] = enc
            struct.pack_into("<H", e, 64, len(enc))
        struct.pack_into("<BBIII", e, 66, typ, 1 if typ else 0, NOSTREAM, NOSTREAM, child)
        struct.pack_into("<IQ", e, 116, start, size)
        return bytes(e)

    directory = (dirent("Root Entry", 5, child=1) + dirent(stream_name, 2, start=2, size=len(payload)) +
                 dirent("", 0) + dirent("", 0))
    data = payload + b"\0" * (n_data * SECT - len(payload))
    return bytes(header) + fat_bytes + directory + data


# --------------------------------------------------------------------------------------------
# Deck builders
# --------------------------------------------------------------------------------------------
def simple_deck(path: Path, n_slides: int, deck_title: str, width=W169, height=H169, notes=False, font=None,
                body_fn=None):
    prs = new_prs(width, height, deck_title)
    for i in range(1, n_slides + 1):
        body = body_fn(i) if body_fn else f"Generated content slide {i} of {n_slides}.\n• Point one\n• Point two"
        s = content_slide(prs, path.name, i, n_slides, f"{deck_title} — slide {i}", body, font=font)
        if notes:
            s.notes_slide.notes_text_frame.text = (
                f"SPEAKER NOTES for slide {i}: presenter view must show this text on the presenter monitor "
                f"and never on the audience screen.")
    save_pptx(prs, path)
    return path


def build_normal():
    simple_deck(OUT / "normal-01-slide.pptx", 1, "Single slide")
    register(name="normal-01-slide.pptx", category="Normal decks",
             tests="Smallest valid 16:9 deck, one slide",
             items="1, 2, 7, 9–16 (baseline)", expected="Launches straight into slideshow on the target monitor; "
             "used for the 100-cycle launch/close and soak runs",
             how="python-pptx")
    simple_deck(OUT / "normal-10-slides-notes.pptx", 10, "Ten slides with notes", notes=True)
    register(name="normal-10-slides-notes.pptx", category="Normal decks",
             tests="10-slide 16:9 deck with speaker notes on every slide",
             items="1, 2, 3, 7", expected="Plays; presenter view ON shows notes on the presenter monitor only; "
             "presenter view OFF shows slides only",
             how="python-pptx")
    simple_deck(OUT / "normal-60-slides.pptx", 60, "Sixty slides")
    register(name="normal-60-slides.pptx", category="Normal decks",
             tests="Long deck (60 slides) — navigation, next/prev/goto, slide index reporting",
             items="1, 7, 10", expected="Plays; agent reports correct current slide after goto(45) and after "
             "agent crash-recovery (item 10)", how="python-pptx")

    prs = new_prs(title="Hidden slide")
    for i in range(1, 6):
        s = content_slide(prs, "hidden-slide.pptx", i, 5, f"Hidden-slide deck — slide {i}",
                          "THIS SLIDE IS HIDDEN — must NOT appear in the show" if i == 3 else "Visible slide")
        if i == 3:
            s._element.set("show", "0")
    save_pptx(prs, OUT / "hidden-slide.pptx")
    register(name="hidden-slide.pptx", category="Normal decks",
             tests="5 slides, slide 3 hidden (p:sld show=\"0\")", items="4, 7",
             expected="Show goes 1→2→4→5; slide 3 never shown; slide count reported as 5 total / 4 visible",
             how="python-pptx + show=\"0\" attribute")


def build_aspects():
    for fname, w, h, label in (("aspect-16x9.pptx", W169, H169, "16:9 (13.333 × 7.5 in)"),
                               ("aspect-4x3.pptx", W43, H43, "4:3 (10 × 7.5 in)"),
                               ("aspect-16x10.pptx", W1610, H1610, "16:10 (12 × 7.5 in)")):
        prs = new_prs(w, h, f"Aspect {label}")
        for i in range(1, 4):
            s = content_slide(prs, fname, i, 3, f"Aspect {label}", "Corner markers must be fully visible")
            for (x, y) in ((0, 0), (w - Inches(0.5), 0), (0, h - Inches(0.5)), (w - Inches(0.5), h - Inches(0.5))):
                box(s, "", x, y, Inches(0.5), Inches(0.5), color="FFFFFF", shape=MSO_SHAPE.RECTANGLE)
        save_pptx(prs, OUT / fname)
        register(name=fname, category="Slide size", tests=f"Slide size {label}; white squares in all four corners",
                 items="4, 7", expected="Fills the display (letter/pillar-boxed as appropriate for the projector); "
                 "all 4 corner markers visible — none cropped", how="python-pptx")


TRANSITIONS = [
    ("Fade", "<p:fade/>"), ("Push up", '<p:push dir="u"/>'), ("Wipe right", '<p:wipe dir="r"/>'),
    ("Split vertical out", '<p:split orient="vert" dir="out"/>'), ("Cover left", '<p:cover dir="l"/>'),
    ("Dissolve", "<p:dissolve/>"), ("Zoom in", '<p:zoom dir="in"/>'), ("Wheel 4 spokes", '<p:wheel spokes="4"/>'),
    ("Circle", "<p:circle/>"), ("Cut through black", '<p:cut thruBlk="1"/>'),
]


def build_transitions():
    prs = new_prs(title="Fade transitions")
    for i in range(1, 6):
        s = content_slide(prs, "transition-fade.pptx", i, 5, "Fade transition", "Every slide fades in (medium)")
        set_transition(s, "<p:fade/>", spd="med")
    save_pptx(prs, OUT / "transition-fade.pptx")
    register(name="transition-fade.pptx", category="Transitions & animations",
             tests="5 slides, fade transition on every slide", items="4",
             expected="Plays with fade transitions identical to opening by hand in PowerPoint", how="python-pptx + "
             "raw <p:transition><p:fade/>", todo="Record side-by-side vs native PowerPoint (fidelity evidence)")

    prs = new_prs(title="Transition variety")
    n = len(TRANSITIONS)
    for i, (label, xml) in enumerate(TRANSITIONS, 1):
        s = content_slide(prs, "transition-variety.pptx", i, n, f"Transition: {label}")
        set_transition(s, xml, spd="slow")
    save_pptx(prs, OUT / "transition-variety.pptx")
    register(name="transition-variety.pptx", category="Transitions & animations",
             tests="10 slides, a different built-in transition per slide (" +
                   ", ".join(t[0] for t in TRANSITIONS) + ")", items="4",
             expected="Each transition renders as in native PowerPoint", how="python-pptx + raw <p:transition>",
             todo="Record side-by-side vs native PowerPoint")

    prs = new_prs(title="Auto-advance loop")
    for i in range(1, 5):
        s = content_slide(prs, "transition-auto-advance-loop.pptx", i, 4, "Auto-advance every 3 s, loops",
                          "No clicks needed")
        set_transition(s, "<p:fade/>", spd="fast", adv_tm=3000)

    def loop_props(entries):
        def fix(data):
            root = etree.fromstring(data)
            show = parse_xml(f'<p:showPr {nsdecls("p", "a")} loop="1" useTimings="1" showNarration="1">'
                             '<p:present/><p:sldAll/><p:penClr><a:prstClr val="red"/></p:penClr></p:showPr>')
            ext = root.find(qn("p:extLst"))
            (ext.addprevious(show) if ext is not None else root.append(show))
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return edit_entry(entries, "ppt/presProps.xml", fix)
    save_pptx(prs, OUT / "transition-auto-advance-loop.pptx", mutate=loop_props)
    register(name="transition-auto-advance-loop.pptx", category="Transitions & animations",
             tests="4 slides auto-advancing every 3 s (advTm) with presentation-level loop (showPr loop=1)",
             items="4, 14", expected="Advances unattended and loops until Esc/stop; agent reports slide changes; "
             "suitable holding/soak content", how="python-pptx + raw advTm + presProps showPr")


def anim_shapes(s, labels, top=Inches(2.0), color="FFC000"):
    shapes = []
    for i, lab in enumerate(labels):
        shapes.append(box(s, lab, Inches(0.8), top + Inches(1.3) * i, Inches(4.5), Inches(1.0), color=color))
    return shapes


def build_animations():
    # Appear on click
    prs = new_prs(title="Appear animations")
    for i in range(1, 3):
        s = content_slide(prs, "anim-entrance-appear.pptx", i, 2, "Entrance: Appear (on click)")
        shp = anim_shapes(s, ["Click 1 — appears", "Click 2 — appears", "Click 3 — appears"])
        set_timing(s, [{"steps": [[("appear", x.shape_id)]]} for x in shp])
    save_pptx(prs, OUT / "anim-entrance-appear.pptx")
    register(name="anim-entrance-appear.pptx", category="Transitions & animations",
             tests="Entrance 'Appear' on click, 3 builds per slide", items="4",
             expected="Each next() reveals one box; after 3 builds next() advances slide; prev() un-builds",
             how="python-pptx + raw <p:timing> (presetID 1 entr)", todo="Verify build order/visuals in PowerPoint")

    # Fade / Fly in, click + with + after previous
    prs = new_prs(title="Fade and fly-in animations")
    s = content_slide(prs, "anim-entrance-fade.pptx", 1, 2, "Entrance: Fade (click, with, after previous)")
    a, b, c = anim_shapes(s, ["Click — fade in", "With previous — fade in", "After previous — fade in"])
    set_timing(s, [{"steps": [[("fade_in", a.shape_id), ("fade_in", b.shape_id)], [("fade_in", c.shape_id)]]}])
    s = content_slide(prs, "anim-entrance-fade.pptx", 2, 2, "Entrance: Fly in from bottom (auto on slide load)")
    a, b = anim_shapes(s, ["Auto — fly in", "After previous — fly in"], color="80DEEA")
    set_timing(s, [{"auto": True, "steps": [[("fly_in", a.shape_id)], [("fly_in", b.shape_id)]]}])
    save_pptx(prs, OUT / "anim-entrance-fade.pptx")
    register(name="anim-entrance-fade.pptx", category="Transitions & animations",
             tests="Entrance 'Fade' with click / with-previous / after-previous chaining; 'Fly In' auto-starting "
                   "on slide load", items="4",
             expected="Slide 1: one click fades in boxes 1+2 together then 3 automatically. Slide 2: both boxes "
             "fly in with no click", how="python-pptx + raw <p:timing> (presetID 10 / 2)",
             todo="Verify timing matches native PowerPoint playback")

    # Motion paths
    prs = new_prs(title="Motion paths")
    s = content_slide(prs, "anim-motion-path.pptx", 1, 1, "Motion paths (on click)")
    a = box(s, "Line →", Inches(0.8), Inches(2.2), Inches(2.2), Inches(1.0), shape=MSO_SHAPE.OVAL)
    b = box(s, "Curve ↘", Inches(0.8), Inches(4.0), Inches(2.2), Inches(1.0), shape=MSO_SHAPE.OVAL, color="A5D6A7")
    set_timing(s, [{"steps": [[("path", a.shape_id)]]}, {"steps": [[("path_curve", b.shape_id)]]}])
    save_pptx(prs, OUT / "anim-motion-path.pptx")
    register(name="anim-motion-path.pptx", category="Transitions & animations",
             tests="Custom motion paths: straight line and cubic-Bezier S-curve, 2 s each, on click", items="4",
             expected="Click 1 moves the orange oval right along a line; click 2 moves the green oval along an "
             "S-curve; both end at the path end (fill=hold)", how="python-pptx + raw <p:animMotion>",
             todo="Verify paths render (PowerPoint shows path preview in Normal view)")

    # Mixed sequence + transitions
    prs = new_prs(title="Mixed animation sequence")
    s = content_slide(prs, "anim-mixed-sequence.pptx", 1, 3, "Entrance → emphasis → exit")
    a, b, c = anim_shapes(s, ["Fade in", "Spin (emphasis)", "Fade out (exit)"])
    set_timing(s, [{"steps": [[("fade_in", a.shape_id)]]}, {"steps": [[("spin", b.shape_id)]]},
                   {"steps": [[("fade_out", c.shape_id)]]}])
    set_transition(s, '<p:push dir="l"/>')
    s = content_slide(prs, "anim-mixed-sequence.pptx", 2, 3, "Auto chain on load: appear → spin → fade out")
    a, b, c = anim_shapes(s, ["Appear", "Spin", "Fade out"], color="CE93D8")
    set_timing(s, [{"auto": True, "steps": [[("appear", a.shape_id)], [("spin", b.shape_id)],
                                            [("fade_out", c.shape_id)]]}])
    set_transition(s, "<p:fade/>")
    s = content_slide(prs, "anim-mixed-sequence.pptx", 3, 3, "Fly in + motion path together")
    a = box(s, "Fly + path", Inches(0.8), Inches(3.0), Inches(3), Inches(1.2))
    b = box(s, "With previous", Inches(0.8), Inches(4.8), Inches(3), Inches(1.2), color="FFAB91")
    set_timing(s, [{"steps": [[("fly_in", a.shape_id), ("fly_in", b.shape_id)], [("path", a.shape_id)]]}])
    set_transition(s, '<p:wipe dir="d"/>')
    save_pptx(prs, OUT / "anim-mixed-sequence.pptx")
    register(name="anim-mixed-sequence.pptx", category="Transitions & animations",
             tests="Entrance, emphasis (spin), exit, auto-start chains, multiple effects on one shape, plus "
                   "push/fade/wipe transitions", items="4",
             expected="All effects play in order as in native PowerPoint; exit leaves the shape hidden",
             how="python-pptx + raw <p:timing>/<p:transition>",
             todo="Primary fidelity deck — record agent vs manual PowerPoint side-by-side")


def movie_slide(prs, deck, title, clip, poster, mime, n=1, total=1):
    s = content_slide(prs, deck, n, total, title)
    W = prs.slide_width
    mv = s.shapes.add_movie(str(clip), W - Inches(7.6), Inches(1.8), Inches(6.4), Inches(3.6),
                            poster_frame_image=str(poster), mime_type=mime)
    return s, mv


def build_video(m):
    specs = [("video-embedded-h264.pptx", "h264", "video/mp4", "H.264/AAC in MP4 (480×270, 4 s)",
              "Plays inline on click (PowerPoint's default for inserted video)", ""),
             ("video-embedded-hevc.pptx", "hevc", "video/mp4", "HEVC/H.265 (hvc1) + AAC in MP4 (480×270, 4 s)",
              "Plays if the Windows 'HEVC Video Extensions' codec is installed; otherwise PowerPoint shows a "
              "media error — the agent must detect/report that, not hang",
              "Record result with and without HEVC Video Extensions on both Windows targets"),
             ("video-embedded-prores.pptx", "prores", "video/quicktime", "Apple ProRes 422 Proxy + PCM in MOV "
              "(480×270, 3 s)", "Windows PowerPoint normally cannot decode ProRes: expect 'media unavailable'; "
              "agent must report unsupported codec (preflight should flag it before the event)",
              "Confirm actual Windows behaviour; decide whether ingest transcodes ProRes")]
    for fname, key, mime, desc, expected, todo in specs:
        prs = new_prs(title=fname)
        movie_slide(prs, fname, f"Embedded video: {key.upper()}", m[key], m[f"{key}_poster"], mime)
        save_pptx(prs, OUT / fname)
        register(name=fname, category="Video", tests=f"Embedded video — {desc}", items="5", expected=expected,
                 how="ffmpeg testsrc2 clip + python-pptx add_movie", todo=todo)

    # Autoplay variant
    prs = new_prs(title="Autoplay video")
    s, mv = movie_slide(prs, "video-embedded-h264-autoplay.pptx", "Embedded H.264 — plays automatically",
                        m["h264"], m["h264_poster"], "video/mp4", 1, 2)
    set_timing(s, [{"auto": True, "steps": [[("play_media", mv.shape_id)]]}], videos=[mv.shape_id], media_ms=4000)
    content_slide(prs, "video-embedded-h264-autoplay.pptx", 2, 2, "After the video")
    save_pptx(prs, OUT / "video-embedded-h264-autoplay.pptx")
    register(name="video-embedded-h264-autoplay.pptx", category="Video",
             tests="Embedded H.264 set to start automatically on slide load (mediacall playFrom)", items="5",
             expected="Video starts without a click when slide 1 shows; audio audible; agent can advance after",
             how="python-pptx add_movie + raw mediacall timing",
             todo="Confirm PowerPoint shows 'Start: Automatically' for this video")

    # Linked video variants
    shutil.copyfile(m["h264"], OUT / "linked-clip-h264.mp4")
    register(name="linked-clip-h264.mp4", category="Video",
             tests="Support file: the target of video-linked-relative.pptx (must sit in the same folder)",
             items="5", expected="n/a (support file)", how="ffmpeg testsrc2, H.264/AAC")
    linked = [("video-linked-relative.pptx", "linked-clip-h264.mp4",
               "Linked (not embedded) H.264 video, relative target in the same folder",
               "Plays when the deck and linked-clip-h264.mp4 are in the same folder; the agent's library sync "
               "must copy the linked file alongside the deck",
               "Confirm PowerPoint resolves the relative link; if not, re-link in PowerPoint on Windows and save"),
              ("video-linked-missing.pptx", "linked-clip-DOES-NOT-EXIST.mp4",
               "Linked video whose target file is deliberately absent",
               "Agent must detect and REPORT missing linked media (preflight before show, and at runtime); "
               "show must not hang on a PowerPoint error dialog", ""),
              ("video-linked-absolute-foreign-path.pptx",
               "file:///C:/Users/speaker/Videos/keynote-intro.mp4",
               "Linked video with an absolute path from the speaker's own PC (realistic missing-media case)",
               "Agent reports missing linked media with the original path in the log", "")]
    for fname, target, desc, expected, todo in linked:
        prs = new_prs(title=fname)
        movie_slide(prs, fname, "Linked video", m["h264"], m["h264_poster"], "video/mp4")

        def mut(entries, target=target):
            entries = externalise(entries, "ppt/slides/slide1.xml", (RT_VIDEO, RT_MEDIA), target)

            def fix_media(data):  # p14:media must use r:link (not r:embed) for linked media
                root = etree.fromstring(data)
                for el in root.iter(f"{{{P14_NS}}}media"):
                    rid = el.attrib.pop(f"{{{R_NS}}}embed")
                    el.set(f"{{{R_NS}}}link", rid)
                return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            return edit_entry(entries, "ppt/slides/slide1.xml", fix_media)
        save_pptx(prs, OUT / fname, mutate=mut)
        register(name=fname, category="Video", tests=desc + f" (Target=\"{target}\", TargetMode=External)",
                 items="5", expected=expected, how="python-pptx add_movie, then video/media relationships rewritten "
                 "to External and the embedded media part removed", todo=todo)

    # Standalone playback-path files
    shutil.copyfile(m["h264_1080"], OUT / "video-h264-1080p.mp4")
    register(name="video-h264-1080p.mp4", category="Non-PowerPoint playback",
             tests="Standalone video file, H.264 High/AAC, 1920×1080, 10 s", items="7",
             expected="Plays full-screen on the target monitor via the agent's video path; returns to holding "
             "screen at end", how="ffmpeg testsrc2 + sine")
    shutil.copyfile(m["hevc_1080"], OUT / "video-hevc-1080p.mp4")
    register(name="video-hevc-1080p.mp4", category="Non-PowerPoint playback",
             tests="Standalone video file, HEVC (hvc1)/AAC, 1920×1080, 10 s", items="5, 7",
             expected="Plays if the player/OS has HEVC support; otherwise agent reports unsupported codec",
             how="ffmpeg libx265")


def build_fonts():
    body = ("The quick brown fox jumps over the lazy dog.\n0123456789 — “quotes” — ÆØÅ éèê ß\n"
            "Line lengths reveal substitution: metrics differ between fonts.")
    simple_deck(OUT / "font-custom-montserrat.pptx", 3, "Custom font: Montserrat", font="Montserrat",
                body_fn=lambda i: body)
    register(name="font-custom-montserrat.pptx", category="Fonts",
             tests="All text set in 'Montserrat' (free, SIL OFL, Google Fonts); font NOT embedded", items="6",
             expected="With Montserrat installed on the room PC: renders in Montserrat. Without it: PowerPoint "
             "substitutes silently — agent/preflight should report the missing font",
             how="python-pptx (a:latin typeface=Montserrat)",
             todo="Run twice on Windows: with Montserrat installed (from fonts.google.com) and without")
    simple_deck(OUT / "font-missing.pptx", 3, "Missing font", font="DXG Missing Font Regular",
                body_fn=lambda i: body)
    register(name="font-missing.pptx", category="Fonts",
             tests="Text set in 'DXG Missing Font Regular', a font that does not exist anywhere", items="6",
             expected="Renders with PowerPoint's substitution (typically the theme/default font); substitution "
             "behaviour documented in the PoC report; preflight lists the font as missing",
             how="python-pptx (a:latin typeface set to a non-existent font)",
             todo="Screenshot the substituted rendering on both Windows targets")
    simple_deck(OUT / "font-mac-only.pptx", 3, "Mac-only font: Helvetica Neue", font="Helvetica Neue",
                body_fn=lambda i: body)
    register(name="font-mac-only.pptx", category="Fonts",
             tests="Text in 'Helvetica Neue' — typical of decks built on a Mac, not present on Windows", items="6",
             expected="Substituted on Windows (usually Arial); line breaks may shift — report as missing font",
             how="python-pptx")
    register(name="font-embedded.pptx", category="Fonts",
             tests="Deck with an EMBEDDED TrueType font (File › Options › Save › Embed fonts)", items="6",
             expected="Renders in the embedded font even though it is not installed", committed=False,
             how="NOT GENERATED — python-pptx cannot embed fonts (.fntdata obfuscation)",
             todo="TODO (Windows): open font-custom-montserrat.pptx in PowerPoint with Montserrat installed, "
                  "enable 'Embed fonts in the file' (all characters), save as font-embedded.pptx, commit, then "
                  "uninstall Montserrat and verify")


def build_legacy_and_pdf():
    # LibreOffice conversions; isolate the LO profile so a running LibreOffice is not disturbed
    profile = (BUILD / "lo-profile").resolve().as_uri()
    conv = BUILD / "lo"
    conv.mkdir(parents=True, exist_ok=True)
    jobs = [("normal-10-slides-notes.pptx", "legacy-10-slides.ppt", "ppt"),
            ("anim-mixed-sequence.pptx", "legacy-animations.ppt", "ppt"),
            ("aspect-4x3.pptx", "legacy-4x3.ppt", "ppt"),
            ("normal-10-slides-notes.pptx", "pdf-10-pages.pdf", "pdf"),
            ("aspect-4x3.pptx", "pdf-4x3-3-pages.pdf", "pdf")]
    for src, dst, fmt in jobs:
        staged = conv / (Path(dst).stem + ".pptx")
        shutil.copyfile(OUT / src, staged)
        run([SOFFICE, f"-env:UserInstallation={profile}", "--headless", "--convert-to", fmt, "--outdir",
             str(conv), str(staged)], timeout=180)
        shutil.move(conv / dst, OUT / dst)
        staged.unlink()
    register(name="legacy-10-slides.ppt", category="Legacy format",
             tests="Legacy binary PowerPoint 97-2003 (.ppt) — 10 slides with notes", items="7",
             expected="Launches and plays like the .pptx (possibly in Compatibility Mode); no conversion prompt "
             "blocks the show", how="LibreOffice --convert-to ppt from normal-10-slides-notes.pptx")
    register(name="legacy-animations.ppt", category="Legacy format",
             tests="Legacy .ppt carrying transitions and animations", items="4, 7",
             expected="Launches; animations/transitions play (LibreOffice export may simplify some effects — "
             "compare to the .pptx original)", how="LibreOffice --convert-to ppt from anim-mixed-sequence.pptx",
             todo="Optionally replace with a .ppt saved by PowerPoint itself (Save As › PowerPoint 97-2003)")
    register(name="legacy-4x3.ppt", category="Legacy format", tests="Legacy .ppt with 4:3 slide size",
             items="7", expected="Launches, 4:3 pillar-boxed on a 16:9 display",
             how="LibreOffice --convert-to ppt from aspect-4x3.pptx")
    register(name="pdf-10-pages.pdf", category="Non-PowerPoint playback", tests="PDF, 10 pages, 16:9",
             items="7", expected="Opens full-screen via the agent's PDF path; next/prev page work; never a bare "
             "desktop", how="LibreOffice --convert-to pdf from normal-10-slides-notes.pptx")
    register(name="pdf-4x3-3-pages.pdf", category="Non-PowerPoint playback", tests="PDF, 3 pages, 4:3 page size",
             items="7", expected="Full-screen, letterboxed correctly",
             how="LibreOffice --convert-to pdf from aspect-4x3.pptx")


def build_macro():
    prs = new_prs(title="Macro-enabled (structural)")
    for i in range(1, 3):
        content_slide(prs, "macro-enabled-structural.pptm", i, 2, "Macro-enabled deck (.pptm)",
                      "Contains a vbaProject.bin part. Macros must never run on a room PC.")
    project = ("ID=\"{00000000-0000-0000-0000-000000000000}\"\r\nName=\"G0_1_Dummy\"\r\n"
               "' Structural placeholder for the DXG G0-1 corpus - contains NO VBA code.\r\n").encode("ascii")
    vba = cfb_single_stream("PROJECT", project)

    def mut(entries):
        def fix_ct(data):
            root = etree.fromstring(data)
            for ov in root.findall(f"{{{CT_NS}}}Override"):
                if ov.get("PartName") == "/ppt/presentation.xml":
                    ov.set("ContentType", CT_PRES_MACRO)
            if not any(d.get("Extension") == "bin" for d in root.findall(f"{{{CT_NS}}}Default")):
                d = etree.Element(f"{{{CT_NS}}}Default", Extension="bin", ContentType=CT_VBA)
                root.insert(0, d)
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

        def fix_rels(data):
            root = etree.fromstring(data)
            etree.SubElement(root, f"{{{PKG_REL_NS}}}Relationship", Id="rIdVba1", Type=RT_VBA,
                             Target="vbaProject.bin")
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        entries = edit_entry(entries, "[Content_Types].xml", fix_ct)
        entries = edit_entry(entries, "ppt/_rels/presentation.xml.rels", fix_rels)
        return entries + [("ppt/vbaProject.bin", vba)]
    save_pptx(prs, OUT / "macro-enabled-structural.pptm", mutate=mut)
    register(name="macro-enabled-structural.pptm", category="Security",
             tests="Macro-enabled package: macroEnabled main content type + vbaProject relationship + a valid "
                   "OLE2 vbaProject.bin with NO VBA code (structural only)", items="8",
             expected="Agent treats .pptm as macro-enabled: macros blocked/disabled per policy (no enable-content "
             "bar interaction, show still runs) or refused per launch rules. PowerPoint may report the VBA "
             "project as damaged — that must also not block the show",
             how="python-pptx + package surgery; vbaProject.bin written by the generator's CFB writer",
             todo="TODO (Windows): make macro-enabled-real.pptm in PowerPoint with a harmless macro "
                  "(Sub OnSlideShowPageChange(): MsgBox \"macro ran\": End Sub) and confirm it does NOT run "
                  "under the agent's Trust Center policy")


def build_external():
    import xlsxwriter
    xlsx = OUT / "linked-data.xlsx"
    wb = xlsxwriter.Workbook(str(xlsx), {"constant_memory": False})
    wb.set_properties({"title": "G0-1 linked data", "author": "DXG G0-1 corpus generator", "created": FIXED_TIME})
    ws = wb.add_worksheet("Data")
    for r, row in enumerate([["Room", "Sessions"], ["Hall A", 12], ["Hall B", 9], ["Room 101", 7]]):
        ws.write_row(r, 0, row)
    wb.close()
    # normalise zip timestamps for determinism
    write_zip(xlsx, read_zip(xlsx))
    register(name="linked-data.xlsx", category="Security", tests="Support file: target of the linked OLE object "
             "in external-linked-ole-and-picture.pptx", items="8", expected="n/a (support file)", how="xlsxwriter")

    img = OUT / "linked-image.png"
    run([FFMPEG, "-y", "-hide_banner", "-loglevel", "error", "-f", "lavfi", "-i",
         "testsrc=size=640x360:rate=1:duration=1", "-frames:v", "1", *BITEXACT, str(img)])
    register(name="linked-image.png", category="Security", tests="Support file: target of the linked picture",
             items="8", expected="n/a (support file)", how="ffmpeg testsrc")

    # Hyperlinks
    prs = new_prs(title="External hyperlinks")
    s = content_slide(prs, "external-hyperlinks.pptx", 1, 2, "External hyperlinks (click only)")
    tb = textbox(s, "Web link: https://example.com", Inches(0.8), Inches(2.0), Inches(8), Inches(0.8), size=28)
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/"
    tb = textbox(s, "Mail link: mailto:test@example.com", Inches(0.8), Inches(2.9), Inches(8), Inches(0.8), size=28)
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = "mailto:test@example.com"
    btn = box(s, "Shape click → https://example.org", Inches(0.8), Inches(4.0), Inches(6), Inches(1.0))
    btn.click_action.hyperlink.address = "https://example.org/"
    doc = box(s, "Shape click → local file (missing)", Inches(0.8), Inches(5.3), Inches(6), Inches(1.0),
              color="EF9A9A")
    doc.click_action.hyperlink.address = "file:///C:/DXG/missing-handout.docx"
    s2 = content_slide(prs, "external-hyperlinks.pptx", 2, 2, "Action: Run program (notepad.exe)")
    run_btn = box(s2, "Click = Run program notepad.exe", Inches(0.8), Inches(2.5), Inches(7), Inches(1.2),
                  color="FF8A80")
    run_btn.click_action.hyperlink.address = "notepad.exe"  # placeholder, rewritten to ppaction://program below
    over = box(s2, "Mouse-over → https://example.net", Inches(0.8), Inches(4.2), Inches(7), Inches(1.2),
               color="B39DDB")
    over.click_action.hyperlink.address = "https://example.net/"

    def mut(entries):
        def fix(data):
            root = etree.fromstring(data)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                  "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            for sp in root.iterfind(".//p:sp", ns):
                name = sp.find("p:nvSpPr/p:cNvPr", ns)
                hl = name.find("a:hlinkClick", ns)
                text = "".join(sp.itertext())
                if hl is None:
                    continue
                if "Run program" in text:
                    hl.set("action", "ppaction://program")
                if "Mouse-over" in text:  # convert click to hover
                    hover = etree.SubElement(name, f"{{{ns['a']}}}hlinkHover")
                    for k, v in hl.attrib.items():
                        hover.set(k, v)
                    name.remove(hl)
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return edit_entry(entries, "ppt/slides/slide2.xml", fix)
    save_pptx(prs, OUT / "external-hyperlinks.pptx", mutate=mut)
    register(name="external-hyperlinks.pptx", category="Security",
             tests="Hyperlinks: web (text + shape click), mailto, missing local file, mouse-over web link, and a "
                   "'Run program' action (ppaction://program → notepad.exe)", items="8",
             expected="Nothing opens automatically during the show — no browser, mail client or program is "
             "launched unless a presenter explicitly clicks; mouse-over link must not fire from pointer drift "
             "(agent hides the cursor); PowerPoint's security prompt for Run program must not block the show",
             how="python-pptx hyperlinks + raw a:hlinkHover / action=ppaction://program",
             todo="Verify in PowerPoint that the Run-program prompt appears and the agent keeps control")

    # Linked OLE object + linked picture
    prs = new_prs(title="Linked OLE and picture")
    s = content_slide(prs, "external-linked-ole-and-picture.pptx", 1, 1, "Linked OLE worksheet + linked picture")
    s.shapes.add_ole_object(str(xlsx), PROG_ID.XLSX, Inches(0.8), Inches(2.0), Inches(3.0), Inches(2.0))
    s.shapes.add_picture(str(img), Inches(4.5), Inches(2.0), Inches(4.0), Inches(2.25))

    def mut2(entries):
        part = "ppt/slides/slide1.xml"
        entries = externalise(entries, part, (RT_PACKAGE,), "linked-data.xlsx", new_type=RT_OLE)

        # picture: switch its blip from embed to link and externalise only that relationship
        def fix_slide(data):
            root = etree.fromstring(data)
            for ole in root.iter(qn("p:oleObj")):
                emb = ole.find(qn("p:embed"))
                link = etree.Element(qn("p:link"))
                link.set("updateAutomatic", "1")
                emb.addprevious(link)
                ole.remove(emb)
            pics = [p for p in root.iter(qn("p:pic")) if p.getparent().tag != qn("p:oleObj")]
            blip = pics[0].find(".//" + qn("a:blip"))
            rid = blip.attrib.pop(qn("r:embed"))
            blip.set(qn("r:link"), rid)
            fix_slide.pic_rid = rid
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        entries = edit_entry(entries, part, fix_slide)

        def fix_rels(data):
            root = etree.fromstring(data)
            for rel in root:
                if rel.get("Id") == fix_slide.pic_rid:
                    rel.set("Target", "linked-image.png")
                    rel.set("TargetMode", "External")
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return drop_unreferenced(edit_entry(entries, rels_path_for(part), fix_rels))
    save_pptx(prs, OUT / "external-linked-ole-and-picture.pptx", mutate=mut2)
    register(name="external-linked-ole-and-picture.pptx", category="Security",
             tests="Linked (not embedded) OLE Excel worksheet with updateAutomatic=1 (target linked-data.xlsx), "
                   "plus a linked picture (a:blip r:link → linked-image.png)", items="5, 8",
             expected="Opening must not block on the 'Update links?' security prompt; links are not refreshed / "
             "Excel is not launched during the show; linked picture shows if present, placeholder if not",
             how="python-pptx add_ole_object/add_picture, then package surgery: p:embed→p:link, relationships "
                 "made External, embedded parts removed",
             todo="Verify PowerPoint's link prompt behaviour; if it rejects the relative link, recreate via "
                  "Insert › Object › Create from file › Link on Windows")


def build_corrupt():
    good = (OUT / "normal-10-slides-notes.pptx").read_bytes()
    (OUT / "corrupt-truncated.pptx").write_bytes(good[: int(len(good) * 0.6)])
    register(name="corrupt-truncated.pptx", category="Corrupted", corrupt=True,
             tests="Valid deck truncated to 60% of its bytes (zip central directory missing)", items="9, 11",
             expected="Must refuse / show holding screen; PowerPoint 'repair' dialog must not be left on screen; "
             "error logged", how="Byte truncation of normal-10-slides-notes.pptx")

    entries = read_zip(OUT / "normal-10-slides-notes.pptx")
    entries = [(n, (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    b'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld>'
                    b'<p:spTree><p:BROKEN attr="unterminated></p:spTree>') if n == "ppt/slides/slide2.xml" else d)
               for n, d in entries]
    write_zip(OUT / "corrupt-broken-xml.pptx", entries)
    register(name="corrupt-broken-xml.pptx", category="Corrupted", corrupt=True,
             tests="Valid zip, but ppt/slides/slide2.xml is malformed XML", items="9, 11",
             expected="PowerPoint offers to repair: agent must not leave the repair dialog on the audience screen; "
             "refuse (holding screen) or accept repair per policy, and report", how="Zip surgery")

    rnd = random.Random(0x6001)
    (OUT / "corrupt-not-a-zip.pptx").write_bytes(rnd.randbytes(64 * 1024))
    register(name="corrupt-not-a-zip.pptx", category="Corrupted", corrupt=True,
             tests="64 KiB of seeded random bytes with a .pptx extension", items="9, 11",
             expected="Refused before PowerPoint is launched (not a zip), holding screen, logged",
             how="random.Random(0x6001).randbytes")
    (OUT / "corrupt-zero-byte.pptx").write_bytes(b"")
    register(name="corrupt-zero-byte.pptx", category="Corrupted", corrupt=True,
             tests="Empty (0-byte) .pptx — e.g. an interrupted upload", items="9, 11, 12",
             expected="Refused before launch; holding screen; logged", how="Empty file")


def build_images(m):
    shutil.copyfile(m["still"], OUT / "image-1920x1080.png")
    register(name="image-1920x1080.png", category="Non-PowerPoint playback", tests="PNG still, 1920×1080",
             items="7", expected="Shown full-screen on the target monitor via the image path",
             how="ffmpeg testsrc2 single frame")
    run([FFMPEG, "-y", "-hide_banner", "-loglevel", "error", "-i", str(m["still"]), "-q:v", "3", *BITEXACT,
         str(OUT / "image-1920x1080.jpg")])
    register(name="image-1920x1080.jpg", category="Non-PowerPoint playback", tests="JPEG still, 1920×1080",
             items="7", expected="Shown full-screen on the target monitor via the image path",
             how="ffmpeg (from the PNG)")


def build_keynote(enabled: bool):
    dst = OUT / "keynote-10-slides.key"
    produced = False
    if enabled and Path("/Applications/Keynote.app").exists():
        src = (OUT / "normal-10-slides-notes.pptx").resolve()
        script = f'''
        tell application "Keynote"
            set theDoc to open POSIX file "{src}"
            delay 2
            save theDoc in POSIX file "{dst.resolve()}"
            close theDoc saving no
        end tell'''
        try:
            if dst.exists():
                shutil.rmtree(dst) if dst.is_dir() else dst.unlink()
            run(["osascript", "-e", script], timeout=120)
            produced = dst.exists()
        except Exception as exc:  # noqa: BLE001
            log(f"  ! Keynote export failed: {exc}")
    elif dst.exists():
        produced = True  # keep a previously produced file
    register(name="keynote-10-slides.key", category="Non-PowerPoint playback", committed=produced,
             tests="Apple Keynote document (10 slides)", items="7",
             expected="On Windows there is no Keynote: agent must either play a pre-converted PDF/PPTX rendition "
             "or refuse with a clear message + holding screen (decide in PoC)",
             how=("Keynote (macOS) opened normal-10-slides-notes.pptx and saved as .key via osascript "
                  "(generate.py --keynote)") if produced else "NOT GENERATED",
             todo="" if produced else "TODO (Mac): run `generate.py --keynote` on a Mac with Keynote, or open "
             "normal-10-slides-notes.pptx in Keynote › File › Save as keynote-10-slides.key in tests/fixtures/g0-1/")


def build_large(m):
    LARGE.mkdir(parents=True, exist_ok=True)
    from PIL import Image
    rnd = random.Random(0x6002)
    prs = new_prs(title="Large deck")
    n = 13  # 13 × ~6.2 MB incompressible noise PNG ≈ 80 MB
    for i in range(1, n + 1):
        s = content_slide(prs, "large/large-deck-80mb.pptx", i, n, f"Large deck — slide {i}")
        p = BUILD / f"noise-{i:02d}.png"
        if not p.exists():
            Image.frombytes("RGB", (1920, 1080), rnd.randbytes(1920 * 1080 * 3)).save(p, compress_level=1)
        else:
            rnd.randbytes(1920 * 1080 * 3)
        s.shapes.add_picture(str(p), Inches(4.3), Inches(1.8), Inches(8.4), Inches(4.725))
    save_pptx(prs, LARGE / "large-deck-80mb.pptx")
    log(f"  large deck: {(LARGE / 'large-deck-80mb.pptx').stat().st_size / 1e6:.1f} MB")


def register_large():
    register(name="large/large-deck-80mb.pptx", category="Normal decks", committed=False,
             tests="Large deck (~80 MB, 13 slides of incompressible 1080p noise images)", items="1, 7, 12",
             expected="Launch-to-first-slide still ≤10 s P95; SHA-256 verification time recorded",
             how="generate.py --large (GITIGNORED — regenerate on the test box)")


# --------------------------------------------------------------------------------------------
# Verification + manifest
# --------------------------------------------------------------------------------------------
def verify() -> list[str]:
    problems = []
    profile = (BUILD / "lo-profile").resolve().as_uri()
    vdir = BUILD / "verify"
    shutil.rmtree(vdir, ignore_errors=True)
    vdir.mkdir(parents=True)
    for e in ENTRIES:
        p = OUT / e.name
        if not p.exists():
            if e.committed:
                problems.append(f"{e.name}: missing")
            continue
        suffix = p.suffix.lower()
        if suffix in (".pptx", ".pptm"):
            try:
                n = len(Presentation(str(p)).slides)
                ok_pptx = True
            except Exception as exc:  # noqa: BLE001
                ok_pptx, n = False, str(exc)[:80]
            if e.corrupt:
                # broken-xml may still parse for slide count (python-pptx loads parts lazily) — LO decides
                pass
            elif not ok_pptx:
                problems.append(f"{e.name}: python-pptx could not open ({n})")
            staged = vdir / (p.stem + suffix)
            shutil.copyfile(p, staged)
            res = subprocess.run([SOFFICE, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                                  "--outdir", str(vdir), str(staged)], capture_output=True, text=True, timeout=240)
            pdf_ok = (vdir / (p.stem + ".pdf")).exists() and (vdir / (p.stem + ".pdf")).stat().st_size > 0
            status = f"python-pptx={'ok (%s slides)' % n if ok_pptx else 'FAIL'}  LibreOffice->pdf={'ok' if pdf_ok else 'FAIL'}"
            if e.corrupt:
                log(f"  [corrupt] {e.name}: {status}  (failure expected)")
            else:
                log(f"  {e.name}: {status}")
                if not pdf_ok:
                    problems.append(f"{e.name}: LibreOffice conversion failed: {res.stderr[-300:]}")
        elif suffix in (".mp4", ".mov"):
            r = subprocess.run([FFPROBE, "-v", "error", "-show_entries", "stream=codec_name,width,height",
                                "-of", "csv=p=0", str(p)], capture_output=True, text=True)
            log(f"  {e.name}: ffprobe {r.stdout.strip().replace(chr(10), ' | ') or 'FAIL'}")
            if r.returncode:
                problems.append(f"{e.name}: ffprobe failed")
        elif suffix == ".ppt":
            staged = vdir / p.name
            shutil.copyfile(p, staged)
            subprocess.run([SOFFICE, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                            "--outdir", str(vdir), str(staged)], capture_output=True, text=True, timeout=240)
            ok = (vdir / (p.stem + ".pdf")).exists()
            log(f"  {e.name}: LibreOffice->pdf={'ok' if ok else 'FAIL'}")
            if not ok:
                problems.append(f"{e.name}: LibreOffice could not read")
        elif suffix == ".pdf":
            ok = p.read_bytes()[:5] == b"%PDF-"
            log(f"  {e.name}: {'ok' if ok else 'FAIL'} PDF header")
            if not ok:
                problems.append(f"{e.name}: bad PDF")
    return problems


def human(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return str(n)


def path_size(p: Path) -> int:
    if p.is_dir():
        return sum(f.stat().st_size for f in p.rglob("*") if f.is_file())
    return p.stat().st_size


def sha256(p: Path) -> str:
    if p.is_dir():
        return ""
    return hashlib.sha256(p.read_bytes()).hexdigest()


def write_manifest():
    js = []
    total = 0
    for e in ENTRIES:
        p = OUT / e.name
        exists = p.exists() and e.committed  # uncommitted files (large/) never affect the manifest
        size = path_size(p) if exists else None
        if exists and e.committed:
            total += size
        js.append({"file": e.name, "category": e.category, "tests": e.tests, "matrixItems": e.items,
                   "expected": e.expected, "howMade": e.how, "todo": e.todo, "corrupt": e.corrupt,
                   "committed": exists, "bytes": size, "sha256": sha256(p) if exists else None})
    lines = [
        "# G0-1 test corpus — manifest",
        "",
        "Generated by `scripts/g0-1-corpus/generate.py` (see its README). **Do not edit by hand** — edit the "
        "`register(...)` calls in the generator and re-run it. Gate: `docs/PHASE0_GATE.md` §G0-1 (matrix items "
        "1–16) and §G0-7; plan item A5 in `docs/poc/G0-1_PLAN.md`. Machine-readable copy with SHA-256 hashes: "
        "`manifest.json`.",
        "",
        f"Committed corpus size: **{human(total)}** "
        f"({sum(1 for e in ENTRIES if e.committed and (OUT / e.name).exists())} files). "
        "`large/` is gitignored and produced with `--large`.",
        "",
        "Matrix item key: 1 launch reliability · 2 presenter view · 3 multi-monitor · 4 animations/transitions · "
        "5 embedded/linked media · 6 fonts · 7 PPT/PPTX/KEY/PDF/video/image paths · 8 macros & external links · "
        "9 PowerPoint crash/failure detection · 10 agent recovery · 11 holding screen · 12 offline restart · "
        "14 soak.",
        "",
        "Linked-media decks resolve their targets **relative to the deck's folder** — copy the whole folder, "
        "not single files, onto the test machine.",
        "",
    ]
    cats: list[str] = []
    for e in ENTRIES:
        if e.category not in cats:
            cats.append(e.category)
    for cat in cats:
        lines += [f"## {cat}", "", "| File | What it tests | Items | Expected behaviour | Size | How made |",
                  "|---|---|---|---|---|---|"]
        for e in (x for x in ENTRIES if x.category == cat):
            p = OUT / e.name
            size = human(path_size(p)) if p.exists() and e.committed else ("not committed" if e.name.startswith(
                "large/") else "—")
            name = f"`{e.name}`"
            if e.todo.startswith("TODO") and not (p.exists() and e.committed):
                name += " **TODO**"
            def code(t):  # keep literal XML tags visible in rendered Markdown
                return re.sub(r"(<[^<>]+>)", r"`\1`", t)
            exp = code(e.expected) + (f" <br>**Check:** {code(e.todo)}" if e.todo else "")
            cells = [name, code(e.tests), e.items, exp, size, code(e.how)]
            lines.append("| " + " | ".join(c.replace("|", "\\|").replace("\n", " ") for c in cells) + " |")
        lines.append("")
    todos = [e for e in ENTRIES if e.todo]
    lines += ["## Checks that need the real application", "",
              "Nothing below can be confirmed on the build Mac (no PowerPoint). Each must be done once on the "
              "Windows PoC box (or a Mac for Keynote) and the result recorded in `docs/poc/ROOM_AGENT_POC.md`.", ""]
    for e in todos:
        lines.append(f"- [ ] `{e.name}` — {e.todo}")
    lines += ["- [ ] All animation/transition decks — they are hand-written OOXML (p:timing / p:transition). Open "
              "each in PowerPoint's Animation Pane and confirm no repair prompt appears and effects are listed as "
              "described. If PowerPoint repairs a file, re-save it from PowerPoint and commit that copy.", ""]
    (OUT / "MANIFEST.md").write_text("\n".join(lines))
    (OUT / "manifest.json").write_text(json.dumps({"generator": "scripts/g0-1-corpus/generate.py",
                                                    "files": js}, indent=2) + "\n")
    return total


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--large", action="store_true", help="also build large/large-deck-80mb.pptx (gitignored)")
    ap.add_argument("--keynote", action="store_true", help="also export a .key via Keynote (macOS only)")
    ap.add_argument("--no-verify", action="store_true", help="skip python-pptx/LibreOffice verification")
    args = ap.parse_args()

    OUT.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(parents=True, exist_ok=True)
    log("media clips…")
    m = build_media()
    log("decks…")
    build_normal()
    build_aspects()
    build_transitions()
    build_animations()
    build_video(m)
    build_fonts()
    build_macro()
    build_external()
    build_corrupt()
    build_images(m)
    log("LibreOffice conversions (.ppt / .pdf)…")
    build_legacy_and_pdf()
    build_keynote(args.keynote)
    register_large()
    if args.large:
        log("large deck…")
        build_large(m)

    problems = []
    if not args.no_verify:
        log("verifying…")
        problems = verify()
    total = write_manifest()
    log(f"\ncommitted corpus: {human(total)} in {OUT.relative_to(REPO)}")
    if total > 25 * 1024 * 1024:
        problems.append(f"committed corpus is {human(total)} (> 25 MB budget)")
    if problems:
        log("\nPROBLEMS:\n  " + "\n  ".join(problems))
        sys.exit(1)
    log("OK")


if __name__ == "__main__":
    main()
